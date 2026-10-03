#!/usr/bin/env python3
"""G3 的**独立参照实现**：用 openpyxl 造 xlsx 并逐格读回。（xref: docs/02-adr/0006-xlsx-协议契约.md §12）

它为什么不在 `packages/xlsx-protocol` 里：三包（engine / xlsx-protocol / pptx-renderer）
**不放非 TS 资产**。差分本体在 `packages/xlsx-protocol/src/xlsx.differential.spec.ts`，因此
**差分就在 `pnpm gate` 的 `test` 步骤里真实运行**——缺 Python 3 或缺 openpyxl 时该步失败，
而不是静默跳过（与 `tools/cpm-reference/` 同一口径）。

它**不是移植，是"另一种写法"**：与 JS 侧只共享[字段契约]（9 列、容差表、诊断码），
不共享任何一行代码——实现与数据来源都不同（openpyxl 自己读写 zip/XML）。

用法：
    python xlsx_reference.py <工作目录>

产出（全部 UTF-8，末尾换行）：
    <工作目录>/py-fixture.xlsx     由 openpyxl 亲手写出的 fixture（**编码形态刻意与 JS 导出不同**）
    <工作目录>/py-read-report.json openpyxl 读 **py-fixture.xlsx** 的事实
    <工作目录>/js-read-report.json openpyxl 读 **js-export.xlsx**（G3 的导出物）的事实

I/O 口径（实测教训）：本机控制台码页会把中文打成乱码，靠 stdout 比对会得到假失败。
因此**一切结果都写文件**，stdout 只输出人读的摘要。
"""

from __future__ import annotations

import json
import re
import sys
import zipfile
from datetime import date, datetime
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font
except ImportError as error:  # pragma: no cover - 环境问题，必须显式失败
    print(f"[xlsx-reference] 缺少 openpyxl：{error}", file=sys.stderr)
    print("[xlsx-reference] 处置：pip install openpyxl==3.1.5（G3 门禁前置）", file=sys.stderr)
    raise SystemExit(3)

HEADERS = ["WBS", "任务名称", "开始", "完成", "工期", "前置任务", "进度", "里程碑", "备注"]
SHEET_NAME = "任务"

# 1900 日期系统：序列号 0 = 1899-12-30
SERIAL_EPOCH = date(1899, 12, 30)


def serial_to_iso(serial: float) -> str:
    whole = int(serial // 1)
    return (SERIAL_EPOCH.fromordinal(SERIAL_EPOCH.toordinal() + whole)).isoformat()


def iso_to_serial(iso: str) -> int:
    parts = [int(piece) for piece in iso.split("-")]
    target = date(parts[0], parts[1], parts[2])
    return target.toordinal() - SERIAL_EPOCH.toordinal()


# ---------------------------------------------------------------- fixture 定义
# **刻意与 JS 导出不同**：日期用「ISO 文本 / 斜杠文本 / 裸序列号」三种形态混写，
# 进度混用分数与百分数点位，里程碑混用 是/否 与 TRUE/FALSE，依赖用全角分号。
PY_FIXTURE: list[dict] = [
    {
        "wbs": "1",
        "name": "阶段一：准备",
        "start": "2026-10-05",  # ISO 文本
        "duration": 4,
        "progress": 0.5,
        "milestone": "否",
        "notes": "含 = 与 + 前缀文本",
    },
    {
        "wbs": "1.1",
        "name": "需求澄清",
        "start": "2026/10/5",  # 斜杠文本
        "end": "2026-10-08",
        "duration": 3,
        "progress": 50,  # 百分数点位
        "milestone": "TRUE",
    },
    {
        "wbs": "1.2",
        "name": "环境搭建",
        "start": 46300,  # 裸序列号 + 日期格式
        "duration": 2,
        "milestone": "是",
    },
    {
        "wbs": "2",
        "name": "阶段二：实现",
        "start": "2026-10-12",
        "duration": 10,
        "progress": "75%",  # 百分号文本
    },
    {
        "wbs": "2.1",
        "name": "渲染内核 🚀",
        "start": "2026-10-12",
        "duration": 5,
        "progress": 0.25,
        "predecessors": "1.2",  # 省略类型 = FS，lag 0
    },
    {
        "wbs": "2.2",
        "name": "导出器",
        "duration": 3,  # 只有工期、没有日期（DM-05 容错）
        "predecessors": "1.2[SS-1]",
    },
    {
        "wbs": "3",
        "name": "里程碑：内审",
        "start": "2026-10-26",
        "duration": 0,
        "milestone": True,
        "predecessors": "2.1[FF]；2.2[SF+1]",  # **全角分号**
    },
    {
        "wbs": "4",
        "name": "无日期任务",
        "notes": "DM-05",
    },
    {
        "wbs": "5",
        "name": "公式日期（有缓存值）",
        "start_formula": ("A1", "2026-10-20"),  # 公式 + 缓存值
    },
    {
        "wbs": "6",
        "name": "公式无缓存值",
        "start_formula": ("TODAY()", None),  # 公式 + 无缓存值
    },
]


def write_fixture(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    for index, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=index, value=header)
        cell.font = Font(bold=True)

    for offset, row in enumerate(PY_FIXTURE):
        excel_row = 2 + offset
        sheet.cell(row=excel_row, column=1, value=row["wbs"])
        sheet.cell(row=excel_row, column=2, value=row["name"])

        if "start" in row:
            cell = sheet.cell(row=excel_row, column=3)
            value = row["start"]
            if isinstance(value, int):
                cell.value = value  # 裸序列号
            else:
                cell.value = value  # 文本（ISO / 斜杠）
            cell.number_format = "yyyy-mm-dd"
        if "end" in row:
            cell = sheet.cell(row=excel_row, column=4, value=row["end"])
            cell.number_format = "yyyy-mm-dd"
        if "duration" in row:
            sheet.cell(row=excel_row, column=5, value=row["duration"])
        if "predecessors" in row:
            sheet.cell(row=excel_row, column=6, value=row["predecessors"])
        if "progress" in row:
            cell = sheet.cell(row=excel_row, column=7, value=row["progress"])
            cell.number_format = "0%"
        if "milestone" in row:
            sheet.cell(row=excel_row, column=8, value=row["milestone"])
        if "notes" in row:
            sheet.cell(row=excel_row, column=9, value=row["notes"])
        if "start_formula" in row:
            formula, cached = row["start_formula"]
            sheet.cell(row=excel_row, column=3, value=f"={formula}")

    # 未识别列：表头不匹配任何规范列（应报 XLSX_UNRECOGNIZED_COLUMN）
    sheet.cell(row=1, column=10, value="负责人")
    sheet.cell(row=2, column=10, value="张三")

    # 缩进**仅作显示**：给 `1.1` 与 `1.2` 加缩进（解析不得依赖它的取值）
    sheet.cell(row=3, column=2).alignment = Alignment(indent=1)
    sheet.cell(row=4, column=2).alignment = Alignment(indent=1)

    workbook.save(path)

    # 公式的缓存值：openpyxl 不计算，直接按 OOXML 补 `<v>`（这是"造脏文件"的正当手法）
    patch_formula_cache(path)


FORMULA_CACHE = {
    # 「公式日期（有缓存值）」那一行的缓存值：2026-10-20
    "C10": str(iso_to_serial("2026-10-20")),
}


def patch_formula_cache(path: Path) -> None:
    """给公式单元格补 `<v>` 缓存值（openpyxl 自己不计算）。

    实测：openpyxl 落盘的是 `<c r="C10"><f>A1</f><v /></c>`（自闭合的空 `<v>`），
    因此正则必须同时容忍"没有 `<v>`"与"空的 `<v />`"两种形态。
    """
    with zipfile.ZipFile(path) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    sheet_path = "xl/worksheets/sheet1.xml"
    xml = parts[sheet_path].decode("utf-8")
    for address, serial in FORMULA_CACHE.items():
        pattern = re.compile(
            rf'(<c r="{address}"[^>]*>)(<f>[^<]*</f>)(?:<v\s*/>|<v>[^<]*</v>)?(</c>)'
        )
        xml = pattern.sub(
            lambda match: f"{match.group(1)}{match.group(2)}<v>{serial}</v>{match.group(3)}", xml
        )
    parts[sheet_path] = xml.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, content in parts.items():
            target.writestr(name, content)


# ---------------------------------------------------------------- 读取（独立实现）


def read_workbook(path: Path) -> dict:
    """openpyxl 的原样读取：不做容差解读，只报**文件里到底是什么**。"""
    workbook = load_workbook(path, data_only=False)
    sheet = workbook[SHEET_NAME]

    headers = {}
    for column in range(1, sheet.max_column + 1):
        value = sheet.cell(row=1, column=column).value
        headers[column] = value
    header_of = {value: column for column, value in headers.items() if isinstance(value, str)}

    rows = []
    for excel_row in range(2, sheet.max_row + 1):
        if all(sheet.cell(row=excel_row, column=column).value in (None, "") for column in range(1, sheet.max_column + 1)):
            continue
        record = {"row": excel_row, "cells": {}}
        for header, column in header_of.items():
            cell = sheet.cell(row=excel_row, column=column)
            record["cells"][header] = describe(cell.value)
        rows.append(record)

    # data_only=True 的独立通道：读**缓存值**（与上面读公式文本形成互证）。
    # 无缓存值的公式在这一通道下读出 `None` —— 与空单元格**同形**，正是 ADR 0006 §5 的口径。
    cached = load_workbook(path, data_only=True)[SHEET_NAME]
    cached_values = {}
    for excel_row in range(2, cached.max_row + 1):
        value = cached.cell(row=excel_row, column=3).value
        # 键用字符串：JSON 对象键必然是字符串，两侧投影才会同形
        cached_values[str(excel_row)] = describe(value)

    return {
        "sheetName": sheet.title,
        "sheetNames": workbook.sheetnames,
        "headers": [headers[column] for column in sorted(headers)],
        "rows": rows,
        "cachedStartValues": cached_values,
        "maxRow": sheet.max_row,
        "maxColumn": sheet.max_column,
    }


def describe(value: object) -> dict:
    """把一个单元格值描述成可 JSON 化的形状（**类型信息必须保留**）。

    `datetime`/`date` 单独成类：openpyxl 会把**带日期格式的数字**解释成日期对象
    （与 ExcelJS 把序列号解释成 `Date` 同源），因此它既不是"文本"也不是"数值"。
    """
    if value is None:
        return {"type": "empty"}
    if isinstance(value, bool):
        return {"type": "boolean", "value": value}
    if isinstance(value, datetime):
        return {"type": "datetime", "value": value.isoformat()}
    if isinstance(value, date):
        return {"type": "datetime", "value": f"{value.isoformat()}T00:00:00"}
    if isinstance(value, (int, float)):
        return {"type": "number", "value": float(value), "integer": float(value).is_integer()}
    if isinstance(value, str):
        if value.startswith("="):
            return {"type": "formula", "value": value}
        return {"type": "text", "value": value}
    return {"type": "other", "value": str(value)}


def is_formula_fact(fact: dict) -> bool:
    return fact.get("type") == "formula"


# ---------------------------------------------------------------- 由事实推导的"JS 侧应当读到什么"
# 这一节是差分的**判据本身**：它按 ADR 0006 §4 的容差表，独立地把 openpyxl 看到的事实翻成文档语义。
# JS 侧导入出的文档必须与它逐字段一致 —— 两侧同时犯同一个错才可能不一致地通过。

CASE_INSENSITIVE_TRUE = {"是", "true", "1", "y", "yes"}
CASE_INSENSITIVE_FALSE = {"否", "false", "0", "n", "no"}


def resolve_date(fact: dict) -> str | None:
    kind = fact["type"]
    if kind == "empty":
        return None
    if kind in ("datetime", "date"):
        return fact["value"][:10]
    if kind == "number":
        serial = fact["value"]
        if serial < 1:
            return "OUT_OF_RANGE"
        return serial_to_iso(serial)
    if kind == "text":
        text = fact["value"].strip()
        iso = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", text)
        slash = re.fullmatch(r"(\d{4})[/.](\d{1,2})[/.](\d{1,2})", text)
        match = iso or slash
        if match is None:
            try:
                number = float(text)
            except ValueError:
                return "UNPARSABLE"
            return resolve_date({"type": "number", "value": number})
        year, month, day = (int(piece) for piece in match.groups())
        return date(year, month, day).isoformat()
    return "UNPARSABLE"


def resolve_duration(fact: dict) -> object:
    if fact["type"] == "empty":
        return None
    if fact["type"] == "number":
        value = fact["value"]
        if not float(value).is_integer():
            return "NOT_INTEGER"
        if value < 0:
            return "NEGATIVE"
        return int(value)
    if fact["type"] == "text":
        try:
            number = float(fact["value"].strip())
        except ValueError:
            return "NOT_INTEGER"
        return resolve_duration({"type": "number", "value": number})
    return "NOT_INTEGER"


def resolve_progress(fact: dict) -> object:
    if fact["type"] == "empty":
        return None
    if fact["type"] == "text":
        text = fact["value"].strip()
        if text.endswith("%"):
            try:
                percent = float(text[:-1])
            except ValueError:
                return "UNPARSABLE"
            return percent / 100 if 0 <= percent <= 100 else "OUT_OF_RANGE"
    if fact["type"] == "number" or fact["type"] == "text":
        try:
            value = float(fact["value"])
        except (TypeError, ValueError):
            return "UNPARSABLE"
        if value <= 1:
            return value
        return value / 100 if 0 <= value <= 100 else "OUT_OF_RANGE"
    return "UNPARSABLE"


def resolve_milestone(fact: dict) -> object:
    if fact["type"] == "empty":
        return False
    if fact["type"] == "boolean":
        return fact["value"]
    if fact["type"] == "number":
        return fact["value"] != 0
    if fact["type"] == "text":
        normalized = fact["value"].strip().lower()
        if normalized in CASE_INSENSITIVE_TRUE:
            return True
        if normalized in CASE_INSENSITIVE_FALSE:
            return False
    return "UNRECOGNIZED"


def derive_document(report: dict) -> list[dict]:
    """按 ADR 0006 §4/§5 的容差表把"openpyxl 看到的事实"翻成**文档语义**。

    这是差分判据的本体：JS 侧导入出的文档必须与它逐字段一致。
    公式单元格走"两通道互证"：`data_only=False` 读公式文本、`data_only=True` 读缓存值，
    因此"公式有缓存值 / 公式无缓存值 / 本来就是空"三类**可区分**。
    """
    rows = []
    for record in report["rows"]:
        cells = record["cells"]
        excel_row = record["row"]
        cached_start = report["cachedStartValues"].get(str(excel_row), {"type": "empty"})
        start_fact = cells.get("开始", {"type": "empty"})

        if is_formula_fact(start_fact):
            # 公式：只读缓存值；无缓存值 → 按值缺失（并与空单元格区分开）
            start = resolve_date(cached_start) if cached_start["type"] != "empty" else "FORMULA_WITHOUT_CACHE"
        else:
            start = resolve_date(start_fact)

        rows.append(
            {
                "row": excel_row,
                "wbs": resolve_text(cells.get("WBS")),
                "name": resolve_text(cells.get("任务名称")),
                "start": start,
                "end": resolve_date(cells.get("完成", {"type": "empty"})),
                "duration": resolve_duration(cells.get("工期", {"type": "empty"})),
                "progress": resolve_progress(cells.get("进度", {"type": "empty"})),
                "milestone": resolve_milestone(cells.get("里程碑", {"type": "empty"})),
                "predecessors": resolve_text(cells.get("前置任务")),
                "unrecognizedColumns": sorted(
                    header
                    for header, fact in cells.items()
                    if header not in HEADERS and fact["type"] != "empty"
                ),
            }
        )
    return rows


def resolve_text(fact: dict | None) -> str:
    if fact is None or fact["type"] == "empty":
        return ""
    return str(fact["value"])


def main() -> int:
    if len(sys.argv) != 2:
        print("用法：python xlsx_reference.py <工作目录>", file=sys.stderr)
        return 2

    workdir = Path(sys.argv[1])
    workdir.mkdir(parents=True, exist_ok=True)

    py_fixture = workdir / "py-fixture.xlsx"
    write_fixture(py_fixture)

    py_report = read_workbook(py_fixture)
    py_report["derived"] = derive_document(py_report)
    (workdir / "py-read-report.json").write_text(
        json.dumps(py_report, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )

    js_export = workdir / "js-export.xlsx"
    if js_export.exists():
        js_report = read_workbook(js_export)
        js_report["derived"] = derive_document(js_report)
        (workdir / "js-read-report.json").write_text(
            json.dumps(js_report, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8"
        )

    print(f"[xlsx-reference] py-fixture.xlsx：{py_report['maxRow']} 行 × {py_report['maxColumn']} 列")
    print(f"[xlsx-reference] 表名={py_report['sheetName']} 表头={py_report['headers']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
